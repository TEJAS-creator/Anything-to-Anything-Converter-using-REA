"""
Data & Spreadsheet Conversion Engine:
CSV <-> XLSX <-> JSON <-> XML <-> Markdown
"""

import os
import csv
import json
import xml.etree.ElementTree as ET
from xml.dom import minidom
import openpyxl


def _load_data(input_path: str, src_format: str) -> list[dict]:
    """Parses input file into list of row dictionaries."""
    fmt = src_format.lower().replace('.', '')
    rows = []

    if fmt == 'csv':
        with open(input_path, mode='r', encoding='utf-8-sig', errors='replace') as f:
            reader = csv.DictReader(f)
            for r in reader:
                rows.append(dict(r))

    elif fmt in ('xlsx', 'xls'):
        wb = openpyxl.load_workbook(input_path, data_only=True)
        ws = wb.active
        headers = []
        for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if r_idx == 0:
                headers = [str(c or f"Col_{i+1}").strip() for i, c in enumerate(row)]
            else:
                if any(c is not None for c in row):
                    row_dict = {}
                    for i, c in enumerate(row):
                        key = headers[i] if i < len(headers) else f"Col_{i+1}"
                        row_dict[key] = "" if c is None else str(c)
                    rows.append(row_dict)

    elif fmt == 'json':
        with open(input_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
            if isinstance(data, list):
                for item in data:
                    if isinstance(item, dict):
                        rows.append(item)
                    else:
                        rows.append({"value": str(item)})
            elif isinstance(data, dict):
                # If wrapped under a key e.g. {"items": [...]} or just one record
                list_keys = [k for k, v in data.items() if isinstance(v, list)]
                if list_keys:
                    for item in data[list_keys[0]]:
                        if isinstance(item, dict):
                            rows.append(item)
                        else:
                            rows.append({"value": str(item)})
                else:
                    rows.append(data)

    elif fmt == 'xml':
        tree = ET.parse(input_path)
        root = tree.getroot()
        for child in root:
            row_dict = {}
            for sub in child:
                row_dict[sub.tag] = sub.text or ""
            if not row_dict:
                row_dict[child.tag] = child.text or ""
            rows.append(row_dict)

    elif fmt in ('md', 'markdown'):
        with open(input_path, 'r', encoding='utf-8') as f:
            lines = [l.strip() for l in f if l.strip().startswith('|')]
            if len(lines) >= 2:
                headers = [h.strip() for h in lines[0].strip('|').split('|')]
                for line in lines[2:]:
                    cols = [c.strip() for c in line.strip('|').split('|')]
                    row_dict = {}
                    for i, h in enumerate(headers):
                        row_dict[h] = cols[i] if i < len(cols) else ""
                    rows.append(row_dict)

    return rows


def _export_data(rows: list[dict], output_path: str, target_format: str) -> str:
    """Exports list of dictionaries to target format."""
    fmt = target_format.lower().replace('.', '')
    
    # Collect all unique headers preserving order
    headers = []
    for r in rows:
        for k in r.keys():
            if k not in headers:
                headers.append(k)
    if not headers:
        headers = ["Data"]

    if fmt == 'csv':
        with open(output_path, 'w', encoding='utf-8', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=headers)
            writer.writeheader()
            for r in rows:
                writer.writerow(r)

    elif fmt == 'xlsx':
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        # Write header
        ws.append(headers)
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        # Write rows
        for r in rows:
            ws.append([r.get(h, "") for h in headers])
        wb.save(output_path)

    elif fmt == 'json':
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(rows, f, indent=2, ensure_ascii=False)

    elif fmt == 'xml':
        root = ET.Element("records")
        for r in rows:
            item = ET.SubElement(root, "record")
            for k, v in r.items():
                safe_tag = "".join(c if c.isalnum() else "_" for c in str(k)).strip("_") or "field"
                if safe_tag[0].isdigit():
                    safe_tag = f"field_{safe_tag}"
                field = ET.SubElement(item, safe_tag)
                field.text = str(v)
        xml_str = ET.tostring(root, encoding="utf-8")
        parsed = minidom.parseString(xml_str)
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(parsed.toprettyxml(indent="  "))

    elif fmt in ('md', 'markdown'):
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write("| " + " | ".join(headers) + " |\n")
            f.write("| " + " | ".join(["---"] * len(headers)) + " |\n")
            for r in rows:
                f.write("| " + " | ".join([str(r.get(h, "")).replace('|', '\\|') for h in headers]) + " |\n")

    return output_path


def convert_data(input_path: str, output_path: str, src_format: str, target_format: str) -> str:
    """Converts data across CSV, XLSX, JSON, XML, Markdown."""
    rows = _load_data(input_path, src_format)
    return _export_data(rows, output_path, target_format)
